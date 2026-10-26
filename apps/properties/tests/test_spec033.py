import io
import zipfile
from decimal import Decimal

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.core.management.base import CommandError
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from PIL import Image

from apps.people import roles
from apps.properties import importers, importing, services
from apps.properties.models import (
    Category,
    ImportJob,
    Property,
    PropertyChange,
)

from .factories import PropertyFactory

pytestmark = pytest.mark.django_db
User = get_user_model()

HEADER = [
    "ref_externa",
    "titulo",
    "categoria",
    "finalidade",
    "cidade",
    "uf",
    "bairro",
    "preco_venda",
    "preco_aluguel",
    "quartos",
    "comodidades",
    "fotos",
]


@pytest.fixture(autouse=True)
def media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path


@pytest.fixture
def admin_client(client):
    client.force_login(User.objects.create_superuser("admin", "a@a.com", "pw"))
    return client


def png():
    buf = io.BytesIO()
    Image.new("RGB", (40, 30), "blue").save(buf, format="PNG")
    return buf.getvalue()


def sheet(rows, header=HEADER):
    book = Workbook()
    ws = book.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def row(ref="E-1", **over):
    base = {
        "ref_externa": ref,
        "titulo": f"Casa {ref}",
        "categoria": "Casa",
        "finalidade": "Venda",
        "cidade": "Campinas",
        "uf": "SP",
        "bairro": "Cambuí",
        "preco_venda": "450.000,00",
        "preco_aluguel": "",
        "quartos": 3,
        "comodidades": "Piscina; Churrasqueira",
        "fotos": "",
    }
    base.update(over)
    return [base[h] for h in HEADER]


def make_zip(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    return out.getvalue()


def upload(data, name):
    return SimpleUploadedFile(name, data)


def make_job(rows, zip_files=None, user=None):
    return importing.create_job(
        ImportJob.Source.SHEET,
        user,
        upload=upload(sheet(rows), "estoque.xlsx"),
        zip_upload=upload(make_zip(zip_files), "fotos.zip") if zip_files is not None else None,
    )


def run(job):
    importing.preview(job)
    importing.confirm(job)
    return importing.run_all(job)


# ---------- leitura ----------


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("450.000,00", "450000.00"),
        ("R$ 1.250,5", "1250.50"),
        ("450000", "450000.00"),
        ("450.000", "450000.00"),
        (450000.0, "450000.00"),
        ("", None),
        (None, None),
    ],
)
def test_parse_decimal(raw, expected):
    assert importers.parse_decimal(raw) == expected


def test_parse_decimal_rejects_garbage():
    with pytest.raises(ValueError, match="inválido"):
        importers.parse_decimal("abc,def,")


def test_parse_sheet_normalizes_values_and_aliases():
    data = sheet(
        [row(finalidade="Venda e aluguel", preco_aluguel="2.500,00", fotos="a.jpg; b.jpg")],
        header=[
            "ref",
            "title",
            "tipo",
            "finalidade",
            "cidade",
            "estado",
            "bairro",
            "preco",
            "aluguel",
            "quartos",
            "comodidades",
            "fotos",
        ],
    )
    [r] = importers.parse_sheet(data, "x.xlsx")
    assert r["external_ref"] == "E-1" and r["category"] == "Casa" and r["state"] == "SP"
    assert r["purpose"] == "BOTH" and r["sale_price"] == "450000.00"
    assert r["rent_price"] == "2500.00" and r["bedrooms"] == 3
    assert r["features"] == ["Piscina", "Churrasqueira"] and r["photos"] == ["a.jpg", "b.jpg"]
    assert r["row"] == 2


def test_parse_sheet_requires_key_columns_and_data():
    with pytest.raises(ValidationError, match="ref_externa"):
        importers.parse_sheet(sheet([["x"]], header=["titulo"]), "x.xlsx")
    with pytest.raises(ValidationError, match="Nenhuma linha"):
        importers.parse_sheet(sheet([]), "x.xlsx")
    with pytest.raises(ValidationError, match="ler a planilha"):
        importers.parse_sheet(b"not a spreadsheet", "x.xlsx")


def test_parse_sheet_skips_blank_rows_and_reads_csv():
    csv_text = (
        "ref_externa;titulo;categoria;cidade;bairro;preco_venda\n"
        "A;Casa A;Casa;Santos;Gonzaga;100000\n"
        ";;;;;\n"
    )
    rows = importers.parse_sheet(csv_text.encode("utf-8-sig"), "x.csv")
    assert [r["external_ref"] for r in rows] == ["A"]


def test_sheet_template_roundtrips():
    book = load_workbook(io.BytesIO(importers.sheet_template()))
    assert [c.value for c in book["Imóveis"][1]][:3] == ["ref_externa", "titulo", "categoria"]
    parsed = importers.parse_sheet(importers.sheet_template(), "modelo.xlsx")
    assert parsed[0]["external_ref"] == "ERP-001"


# ---------- prévia e validação ----------


def test_preview_counts_and_row_errors():
    PropertyFactory(external_ref="E-2")
    job = make_job(
        [
            row("E-1"),
            row("E-2"),
            row("E-1"),  # repetida
            row("E-3", categoria="Inexistente"),
            row("E-4", preco_venda=""),
            row("E-5", finalidade="Aluguel", preco_aluguel=""),
            row("", titulo=""),
            row("E-7", cidade=""),
        ]
    )
    info = importing.preview(job)
    assert (info["new"], info["updated"]) == (1, 1)
    messages = {i["row"]: i["message"] for i in info["invalid"]}
    assert "repetida" in messages[4]
    assert "categoria não encontrada" in messages[5]
    assert "preço de venda" in messages[6]
    assert "valor do aluguel" in messages[7]
    assert "referência externa vazia" in messages[8] and "título vazio" in messages[8]
    assert "cidade e bairro" in messages[9]


def test_create_job_rejects_bad_inputs():
    with pytest.raises(ValidationError, match="planilha"):
        importing.create_job(ImportJob.Source.SHEET, None)
    with pytest.raises(ValidationError, match="XML"):
        importing.create_job(ImportJob.Source.XML, None)
    with pytest.raises(ValidationError, match="ZIP"):
        importing.create_job(
            ImportJob.Source.SHEET,
            None,
            upload=upload(sheet([row()]), "a.xlsx"),
            zip_upload=upload(b"not a zip", "f.zip"),
        )
    assert not ImportJob.objects.exists()


# ---------- importação ----------


def test_import_creates_properties_photos_history():
    job = make_job(
        [row("E-1", fotos="sala.png"), row("E-2")],
        zip_files={"x/sala.png": png(), "E-2/1.png": png(), "E-2/2.png": png()},
    )
    run(job)
    job.refresh_from_db()
    assert job.status == ImportJob.Status.DONE
    assert (job.created_count, job.updated_count, job.error_count) == (2, 0, 0)
    first = Property.objects.get(external_ref="E-1")
    second = Property.objects.get(external_ref="E-2")
    assert first.photos.count() == 1 and second.photos.count() == 2
    assert second.photos.first().is_cover
    assert str(first.sale_price) == "450000.00" and first.status == Property.Status.DRAFT
    assert sorted(f.name for f in first.features.all()) == ["Churrasqueira", "Piscina"]
    rows = PropertyChange.objects.filter(property=first)
    assert {r.source for r in rows} == {f"Importação #{job.pk}"}
    assert rows.count() == 2  # cadastro + fotos (uma linha para todas as fotos)


def test_reimport_updates_without_duplicating_or_readding_photos():
    zipped = {"E-1/a.png": png()}
    run(make_job([row("E-1")], zipped))
    prop = Property.objects.get(external_ref="E-1")
    job = run(make_job([row("E-1", titulo="Novo título", preco_venda="500.000,00")], zipped))
    assert (job.created_count, job.updated_count) == (0, 1)
    prop.refresh_from_db()
    assert prop.title == "Novo título" and prop.sale_price == Decimal("500000.00")
    assert Property.objects.count() == 1 and prop.photos.count() == 1


def test_invalid_rows_are_reported_and_do_not_stop_the_rest():
    job = run(make_job([row("E-1"), row("E-2", categoria="Nada"), row("E-3")]))
    assert (job.created_count, job.error_count) == (2, 1)
    [issue] = job.errors
    assert issue["row"] == 3 and issue["level"] == "error"
    assert not Property.objects.filter(external_ref="E-2").exists()


def test_missing_and_invalid_photos_are_warnings_not_errors():
    job = run(
        make_job(
            [row("E-1", fotos="nao-existe.png; ruim.png")],
            zip_files={"ruim.png": b"nao sou imagem", "outra.png": png()},
        )
    )
    assert (job.created_count, job.error_count) == (1, 0)
    levels = {(i["level"]) for i in job.errors}
    messages = " ".join(i["message"] for i in job.errors)
    assert levels == {"warning"} and "nao-existe.png" in messages and "ruim.png" in messages


def test_photos_without_zip_warn(monkeypatch):
    job = run(make_job([row("E-1", fotos="a.png")]))
    assert job.created_count == 1 and "nenhum ZIP" in job.errors[0]["message"]


def test_photo_urls_are_downloaded_with_the_safe_fetcher(monkeypatch):
    calls = []
    monkeypatch.setattr(services, "fetch_image_bytes", lambda url: calls.append(url) or png())
    job = run(make_job([row("E-1", fotos="https://cdn.exemplo.com/a.png")]))
    assert calls == ["https://cdn.exemplo.com/a.png"]
    assert Property.objects.get(external_ref="E-1").photos.count() == 1
    assert job.error_count == 0


def test_private_photo_url_is_a_warning():
    job = run(make_job([row("E-1", fotos="http://127.0.0.1/a.png")]))
    assert job.created_count == 1 and job.errors[0]["level"] == "warning"


def test_batches_resume_where_they_stopped():
    job = make_job([row(f"E-{i}") for i in range(25)])
    importing.preview(job)
    importing.confirm(job)
    importing.process_batch(job, max_rows=10)
    assert (job.processed, job.status) == (10, ImportJob.Status.RUNNING)
    fresh = ImportJob.objects.get(pk=job.pk)
    importing.process_batch(fresh, max_rows=10)
    importing.process_batch(fresh, max_rows=10)
    assert fresh.processed == 25 and fresh.status == ImportJob.Status.DONE
    assert fresh.finished_at and Property.objects.count() == 25


def test_time_budget_ends_the_batch_early():
    job = make_job([row(f"E-{i}") for i in range(5)])
    importing.preview(job)
    importing.confirm(job)
    importing.process_batch(job, max_rows=20, budget=0)
    assert job.processed == 1


def test_confirm_only_from_preview_and_not_with_unmapped_types():
    job = make_job([row()])
    importing.confirm(job)
    with pytest.raises(ValidationError):
        importing.confirm(job)


# ---------- ZIP ----------


def test_zip_reads_by_name_and_folder_without_extracting():
    z = importers.PhotoArchive(
        io.BytesIO(
            make_zip({"E-1/b.png": png(), "E-1/a.png": png(), "solta.PNG": png(), "../x.txt": b"x"})
        )
    )
    found, missing = z.photos_for("E-1", [])
    assert [n for n, _ in found] == ["E-1/a.png", "E-1/b.png"] and missing == []
    found, missing = z.photos_for("E-9", ["SOLTA.png", "falta.png"])
    assert [n for n, _ in found] == ["solta.PNG"] and missing == ["falta.png"]


def test_zip_too_many_members_is_refused(monkeypatch):
    monkeypatch.setattr(importers, "MAX_ZIP_MEMBERS", 1)
    with pytest.raises(ValidationError, match="grande demais"):
        importers.PhotoArchive(io.BytesIO(make_zip({"a.png": png(), "b.png": png()})))


# ---------- XML ----------

XML = """<?xml version="1.0" encoding="UTF-8"?>
<ListingDataFeed>
  <Listings>
    <Listing>
      <ListingID>V-1</ListingID>
      <Title>Apartamento Gonzaga</Title>
      <TransactionType>For Sale</TransactionType>
      <Details>
        <PropertyType>Residential / Apartment</PropertyType>
        <Description>Vista mar</Description>
        <ListPrice currency="BRL">750000</ListPrice>
        <PropertyAdministrationFee>850</PropertyAdministrationFee>
        <LivingArea unit="square metres">90</LivingArea>
        <Bedrooms>3</Bedrooms><Bathrooms>2</Bathrooms><Garage>1</Garage>
      </Details>
      <Features><Feature>Piscina</Feature><Feature>Sauna</Feature></Features>
      <Location>
        <Country abbreviation="BR"><State abbreviation="SP"><City>Santos</City>
        <Neighborhood>Gonzaga</Neighborhood><Address>Av. Ana Costa</Address>
        <StreetNumber>10</StreetNumber><PostalCode>11060-001</PostalCode></State></Country>
      </Location>
      <Media><Item medium="image">https://cdn.exemplo.com/1.png</Item>
             <Item medium="video">https://cdn.exemplo.com/v.mp4</Item></Media>
    </Listing>
    <Listing>
      <ListingID>V-2</ListingID>
      <Title>Sítio</Title>
      <TransactionType>For Rent</TransactionType>
      <Details><PropertyType>Rural / Farm</PropertyType><RentalPrice>3000</RentalPrice></Details>
      <Location><State abbreviation="SP"><City>Atibaia</City>
        <Neighborhood>Rural</Neighborhood></State></Location>
    </Listing>
  </Listings>
</ListingDataFeed>"""


def xml_job(xml=XML):
    return importing.create_job(ImportJob.Source.XML, None, upload=upload(xml.encode(), "feed.xml"))


def test_parse_vrsync_fields():
    rows = importers.parse_vrsync(XML.encode())
    first, second = rows
    assert first["external_ref"] == "V-1" and first["purpose"] == "SALE"
    assert first["sale_price"] == "750000.00" and first["condo_fee"] == "850.00"
    assert first["area_built"] == "90.00" and first["parking_spaces"] == 1
    assert (first["city"], first["state"], first["neighborhood"]) == ("Santos", "SP", "Gonzaga")
    assert first["features"] == ["Piscina", "Sauna"]
    assert first["photos"] == ["https://cdn.exemplo.com/1.png"]  # vídeo fica de fora
    assert first["portal_type"] == "Residential / Apartment"
    assert second["purpose"] == "RENT" and second["rent_price"] == "3000.00"


def test_xml_types_are_suggested_and_unmapped_ones_block_confirmation():
    job = xml_job()
    assert job.type_map == {"Residential / Apartment": "Apartamento", "Rural / Farm": "Fazenda"}
    job = xml_job(XML.replace("Rural / Farm", "Algo / Estranho"))
    assert importing.unmapped_types(job) == ["Algo / Estranho"]
    with pytest.raises(ValidationError):
        importing.confirm(job)
    importing.set_type_map(job, {"Algo / Estranho": "Casa"})
    assert importing.unmapped_types(job) == []


def test_xml_import_with_mapping_and_ignored_type(monkeypatch):
    monkeypatch.setattr(services, "fetch_image_bytes", lambda url: png())
    job = xml_job()
    importing.set_type_map(job, {"Rural / Farm": ""})  # ignorar sítios
    info = importing.preview(job)
    assert (info["new"], info["skipped"], info["invalid"]) == (1, 1, [])
    importing.confirm(job)
    importing.run_all(job)
    prop = Property.objects.get()
    assert prop.external_ref == "V-1" and prop.category.name == "Apartamento"
    assert prop.neighborhood.city.name == "Santos" and prop.photos.count() == 1
    assert job.created_count == 1 and job.error_count == 0


def test_xml_rejects_entities_and_non_feeds():
    bomb = '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa"><!ENTITY b "&a;&a;&a;">]><x>&b;</x>'
    with pytest.raises(ValidationError, match="XML"):
        importers.parse_vrsync(bomb.encode())
    xxe = '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>'
    with pytest.raises(ValidationError, match="XML"):
        importers.parse_vrsync(xxe.encode())
    with pytest.raises(ValidationError, match="Listing"):
        importers.parse_vrsync(b"<a><b/></a>")
    with pytest.raises(ValidationError, match="XML"):
        importers.parse_vrsync(b"not xml")


def test_xml_feed_url_uses_the_safe_fetcher(monkeypatch):
    monkeypatch.setattr(services, "fetch_bytes", lambda url, *a, **k: XML.encode())
    job = importing.create_job(
        ImportJob.Source.XML, None, feed_url="https://portal.exemplo.com/feed.xml"
    )
    assert job.total == 2 and job.feed_url.endswith("feed.xml")


def test_xml_feed_url_to_private_host_is_refused():
    with pytest.raises(ValidationError):
        importing.create_job(ImportJob.Source.XML, None, feed_url="http://127.0.0.1/feed.xml")
    assert not ImportJob.objects.exists()


# ---------- painel ----------


def test_admin_upload_preview_confirm_and_progress(admin_client):
    add = reverse("admin:properties_importjob_add")
    page = admin_client.get(add).content.decode()
    assert "modelo (XLSX)" in page and reverse("admin:properties_importjob_template") in page

    r = admin_client.post(
        add,
        {
            "source": "SHEET",
            "file": upload(sheet([row("E-1"), row("E-2", categoria="Nada")]), "e.xlsx"),
            "feed_url": "",
        },
    )
    job = ImportJob.objects.get()
    page_url = reverse("admin:properties_importjob_change", args=[job.pk])
    assert r.status_code == 302 and r.url == page_url
    html = admin_client.get(page_url).content.decode()
    assert "Prévia" in html and "categoria não encontrada" in html
    assert "Confirmar e importar 1 imóvel" in html

    admin_client.post(page_url, {"action": "confirm"})
    job.refresh_from_db()
    assert job.status == ImportJob.Status.RUNNING
    assert "Importando 0/2" in admin_client.get(page_url).content.decode()

    step = reverse("admin:properties_importjob_step", args=[job.pk])
    r = admin_client.post(step, HTTP_HX_REQUEST="true")
    html = r.content.decode()
    assert r.status_code == 200 and "Importação concluída: 2/2" in html
    assert "hx-post" not in html and "Linha 3" in html
    assert Property.objects.filter(external_ref="E-1").exists()


def test_admin_non_htmx_step_redirects_and_xml_mapping_flow(admin_client):
    xml = XML.replace("Residential / Apartment", "Residential / Estranho")
    admin_client.post(
        reverse("admin:properties_importjob_add"),
        {"source": "XML", "file": upload(xml.encode(), "feed.xml"), "feed_url": ""},
    )
    job = ImportJob.objects.get()
    url = reverse("admin:properties_importjob_change", args=[job.pk])
    assert "Ligue os tipos do portal" in admin_client.get(url).content.decode()
    admin_client.post(url, {"action": "map", "map__Residential / Estranho": "Casa"})
    job.refresh_from_db()
    assert job.type_map["Residential / Estranho"] == "Casa"
    admin_client.post(url, {"action": "map", "map__Residential / Estranho": "Categoria Falsa"})
    job.refresh_from_db()
    assert job.type_map["Residential / Estranho"] == "Casa"  # categoria inválida é ignorada
    html = admin_client.get(url).content.decode()
    assert "Prévia" in html and "Ligue os tipos" not in html


def test_admin_rejects_invalid_upload(admin_client):
    r = admin_client.post(
        reverse("admin:properties_importjob_add"),
        {"source": "SHEET", "file": upload(b"lixo", "e.xlsx"), "feed_url": ""},
    )
    assert r.status_code == 200 and "ler a planilha" in r.content.decode()
    assert not ImportJob.objects.exists()


def test_admin_template_download(admin_client):
    r = admin_client.get(reverse("admin:properties_importjob_template"))
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
    assert "modelo-importacao-imoveis.xlsx" in r["Content-Disposition"]
    assert load_workbook(io.BytesIO(r.content))["Imóveis"]["A1"].value == "ref_externa"


@pytest.mark.parametrize(
    ("role", "allowed"),
    [(roles.ADMIN, True), (roles.MANAGER, True), (roles.BROKER, False), (roles.CLERK, False)],
)
def test_only_admin_and_manager_can_import(client, role, allowed):
    user = User.objects.create_user("u", password="pw", is_staff=True)
    roles.apply_role(user, role)
    client.force_login(user)
    job = make_job([row()])
    for url in (
        reverse("admin:properties_importjob_add"),
        reverse("admin:properties_importjob_template"),
        reverse("admin:properties_importjob_change", args=[job.pk]),
    ):
        assert (client.get(url).status_code == 200) is allowed
    assert client.post(reverse("admin:properties_importjob_step", args=[job.pk])).status_code in (
        (302,) if allowed else (403,)
    )


def test_sidebar_item_follows_permission(client):
    from apps.core.admin_navigation import can_add_importjob

    manager = User.objects.create_user("m", password="pw", is_staff=True)
    roles.apply_role(manager, roles.MANAGER)
    broker = User.objects.create_user("b", password="pw", is_staff=True)
    roles.apply_role(broker, roles.BROKER)
    request = type("R", (), {})()
    request.user = manager
    assert can_add_importjob(request)
    request.user = User.objects.get(pk=broker.pk)
    assert not can_add_importjob(request)


# ---------- comando ----------


def test_command_imports_sheet_with_zip(tmp_path, capsys):
    xlsx = tmp_path / "e.xlsx"
    xlsx.write_bytes(sheet([row("E-1"), row("E-2", categoria="Nada")]))
    zipped = tmp_path / "f.zip"
    zipped.write_bytes(make_zip({"E-1/a.png": png()}))
    call_command("import_properties", str(xlsx), "--zip", str(zipped))
    out = capsys.readouterr().out
    assert "1 criado(s)" in out and "1 com erro" in out
    assert Property.objects.get(external_ref="E-1").photos.count() == 1


def test_command_dry_run_imports_nothing(tmp_path, capsys):
    xlsx = tmp_path / "e.xlsx"
    xlsx.write_bytes(sheet([row("E-1")]))
    call_command("import_properties", str(xlsx), "--dry-run")
    assert "1 novo(s)" in capsys.readouterr().out
    assert not Property.objects.exists() and not ImportJob.objects.exists()


def test_command_xml_needs_mapping_or_options(tmp_path):
    feed = tmp_path / "f.xml"
    feed.write_text(XML.replace("Residential / Apartment", "Algo / Estranho"), encoding="utf-8")
    with pytest.raises(CommandError, match="Algo / Estranho"):
        call_command("import_properties", str(feed))
    call_command(
        "import_properties", str(feed), "--map", "Algo / Estranho=Casa", "--map", "Rural / Farm="
    )
    assert Property.objects.get().category == Category.objects.get(name="Casa")


def test_command_missing_file():
    with pytest.raises(CommandError, match="não encontrado"):
        call_command("import_properties", "/nao/existe.xlsx")
