"""Leitura de arquivos de importação (planilha/CSV e XML VivaReal/ZAP) e validação das linhas.

Toda fonte vira a mesma lista de linhas (dicionários com os campos de `services.upsert_property`
mais `row` (número), `photos` e, no XML, `portal_type`), guardada em `ImportJob.rows`.
"""

import csv
import io
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation

import defusedxml.ElementTree as ET
from django.core.exceptions import ValidationError
from openpyxl import Workbook, load_workbook

from .models import Category, Property

MAX_SHEET_BYTES = 20 * 1024 * 1024
MAX_XML_BYTES = 50 * 1024 * 1024
MAX_ROWS = 5000
MAX_ZIP_MEMBERS = 20000
MAX_ZIP_UNCOMPRESSED = 2 * 1024 * 1024 * 1024
IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".gif")

# Colunas do modelo de planilha: (cabeçalho, exemplo, aliases aceitos na leitura).
COLUMNS = [
    ("ref_externa", "ERP-001", ["referencia", "ref", "codigo_externo", "id"]),
    ("titulo", "Casa 3 quartos com piscina", ["title"]),
    ("categoria", "Casa", ["tipo", "category"]),
    ("finalidade", "Venda", ["purpose"]),
    ("situacao", "Rascunho", ["status"]),
    ("cidade", "Campinas", ["city"]),
    ("uf", "SP", ["estado", "state"]),
    ("bairro", "Cambuí", ["neighborhood"]),
    ("rua", "Rua das Flores", ["street", "logradouro"]),
    ("numero", "120", ["number"]),
    ("complemento", "Apto 12", ["complement"]),
    ("cep", "13000-000", ["zip_code"]),
    ("preco_venda", "450000", ["sale_price", "preco"]),
    ("preco_aluguel", "", ["rent_price", "aluguel"]),
    ("condominio", "", ["condo_fee"]),
    ("iptu", "", []),
    ("quartos", "3", ["bedrooms"]),
    ("suites", "1", []),
    ("banheiros", "2", ["bathrooms"]),
    ("vagas", "2", ["parking_spaces"]),
    ("area_total", "300", []),
    ("area_construida", "180", ["area_built"]),
    ("descricao", "Casa ampla, quintal e churrasqueira.", ["description"]),
    ("comodidades", "Piscina; Churrasqueira", ["features"]),
    ("fotos", "ERP-001/fachada.jpg; sala.jpg", ["photos", "imagens"]),
]
FIELD_OF = {
    "ref_externa": "external_ref",
    "titulo": "title",
    "categoria": "category",
    "finalidade": "purpose",
    "situacao": "status",
    "cidade": "city",
    "uf": "state",
    "bairro": "neighborhood",
    "rua": "street",
    "numero": "number",
    "complemento": "complement",
    "cep": "zip_code",
    "preco_venda": "sale_price",
    "preco_aluguel": "rent_price",
    "condominio": "condo_fee",
    "iptu": "iptu",
    "quartos": "bedrooms",
    "suites": "suites",
    "banheiros": "bathrooms",
    "vagas": "parking_spaces",
    "area_total": "area_total",
    "area_construida": "area_built",
    "descricao": "description",
    "comodidades": "features",
    "fotos": "photos",
}
MONEY = {"sale_price", "rent_price", "condo_fee", "iptu", "area_total", "area_built"}
INTEGERS = {"bedrooms", "suites", "bathrooms", "parking_spaces"}

PURPOSES = {
    "venda": "SALE",
    "vender": "SALE",
    "comprar": "SALE",
    "sale": "SALE",
    "aluguel": "RENT",
    "locacao": "RENT",
    "alugar": "RENT",
    "rent": "RENT",
    "venda e aluguel": "BOTH",
    "venda/aluguel": "BOTH",
    "ambos": "BOTH",
    "both": "BOTH",
}


def _plain(text):
    text = unicodedata.normalize("NFKD", str(text or ""))
    return "".join(c for c in text if not unicodedata.combining(c)).strip().lower()


def _header_map():
    mapping = {}
    for header, _example, aliases in COLUMNS:
        for name in (header, *aliases):
            mapping[_plain(name).replace(" ", "_")] = header
    return mapping


def parse_decimal(value):
    """ "450.000,00", "450000" e números do Excel viram "450000.00"; vazio vira None."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return f"{Decimal(str(value)):.2f}"
    text = re.sub(r"[^\d,.\-]", "", str(value))
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif text.count(".") > 1 or re.fullmatch(r"-?\d{1,3}\.\d{3}", text):
        text = text.replace(".", "")
    try:
        return f"{Decimal(text):.2f}"
    except InvalidOperation as exc:
        raise ValueError(f"valor numérico inválido: {value!r}") from exc


def parse_int(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(Decimal(str(value).replace(",", ".")))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"número inválido: {value!r}") from exc


def split_list(value):
    return [p.strip() for p in re.split(r"[;\n|]", str(value or "")) if p.strip()]


def _clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# ---------- planilha ----------


def _read_table(data, filename):
    """Linhas da planilha como listas de valores (primeira linha = cabeçalho)."""
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else None
        return [list(r) for r in csv.reader(io.StringIO(text), dialect or csv.excel)]
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.active
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    except Exception as exc:  # noqa: BLE001 - arquivo corrompido ou não é XLSX
        raise ValidationError(
            "Não foi possível ler a planilha. Use o modelo em XLSX ou um CSV."
        ) from exc


def parse_sheet(data, filename):
    """Devolve a lista de linhas normalizadas. Levanta ValidationError se faltar coluna-chave."""
    if len(data) > MAX_SHEET_BYTES:
        raise ValidationError("Planilha maior que 20 MB.")
    table = _read_table(data, filename)
    if not table:
        raise ValidationError("A planilha está vazia.")
    aliases = _header_map()
    header = [aliases.get(_plain(h).replace(" ", "_")) for h in table[0]]
    missing = [c for c in ("ref_externa", "titulo", "categoria") if c not in header]
    if missing:
        raise ValidationError(
            f"Colunas obrigatórias ausentes: {', '.join(missing)}. Baixe o modelo de planilha."
        )
    rows = []
    for number, cells in enumerate(table[1:], start=2):
        if not any(_clean_cell(c) for c in cells):
            continue
        raw = {
            col: _clean_cell(cells[i]) if i < len(cells) else ""
            for i, col in enumerate(header)
            if col
        }
        rows.append(_sheet_row(number, raw))
        if len(rows) > MAX_ROWS:
            raise ValidationError(f"Mais de {MAX_ROWS} linhas: divida em arquivos menores.")
    if not rows:
        raise ValidationError("Nenhuma linha de dados na planilha.")
    return rows


def _sheet_row(number, raw):
    row = {"row": number, "parse_errors": [], "photos": split_list(raw.get("fotos"))}
    for column, field in FIELD_OF.items():
        value = raw.get(column, "")
        if field == "photos" or value == "":
            continue
        try:
            if field in MONEY:
                row[field] = parse_decimal(value)
            elif field in INTEGERS:
                row[field] = parse_int(value)
            elif field == "features":
                row[field] = split_list(value)
            elif field == "purpose":
                row[field] = PURPOSES.get(_plain(value), value.upper())
            elif field == "status":
                row[field] = _status(value)
            elif field == "state":
                row[field] = value.upper()
            else:
                row[field] = value
        except ValueError as exc:
            row["parse_errors"].append(f"{column}: {exc}")
    return row


def _status(value):
    names = {_plain(label): code for code, label in Property.Status.choices}
    return names.get(_plain(value), str(value).upper())


# ---------- XML (VRSync: VivaReal/ZAP) ----------

TYPE_HINTS = {
    "apartment": "Apartamento",
    "home": "Casa",
    "house": "Casa",
    "condo": "Casa",
    "penthouse": "Cobertura",
    "land": "Terreno",
    "lot": "Terreno",
    "office": "Comercial",
    "commercial": "Comercial",
    "farm": "Fazenda",
}


def suggest_category(portal_type):
    last = _plain(re.split(r"[/>]", portal_type)[-1])
    for hint, name in TYPE_HINTS.items():
        if hint in last:
            return Category.objects.filter(name__iexact=name).values_list("name", flat=True).first()
    return None


def _text(node, path):
    found = node.find(path) if node is not None else None
    return (found.text or "").strip() if found is not None and found.text else ""


def _strip_ns(root):
    for el in root.iter():
        if isinstance(el.tag, str) and "}" in el.tag:
            el.tag = el.tag.split("}", 1)[1]
    return root


def parse_vrsync(data):
    """Linhas do feed VivaReal/ZAP (VRSync). Cada imóvel traz `portal_type` para o mapeamento."""
    if len(data) > MAX_XML_BYTES:
        raise ValidationError("XML maior que 50 MB.")
    try:
        root = _strip_ns(ET.fromstring(data))
    except Exception as exc:  # noqa: BLE001 - XML inválido ou bloqueado pelo defusedxml
        raise ValidationError(
            "Não foi possível ler o XML (arquivo inválido ou não permitido)."
        ) from exc
    listings = root.findall(".//Listing")
    if not listings:
        raise ValidationError(
            "Nenhum <Listing> encontrado: o arquivo não parece um feed VivaReal/ZAP."
        )
    if len(listings) > MAX_ROWS:
        raise ValidationError(f"Mais de {MAX_ROWS} imóveis no feed: divida em arquivos menores.")
    return [_listing_row(i, node) for i, node in enumerate(listings, start=1)]


def _listing_row(number, node):
    row = {"row": number, "parse_errors": [], "photos": []}
    details = node.find("Details")
    transaction = _plain(_text(node, "TransactionType"))
    purpose = "BOTH" if "sale" in transaction and "rent" in transaction else None
    purpose = purpose or ("RENT" if "rent" in transaction else "SALE")
    location = node.find("Location")
    state = location.find(".//State") if location is not None else None
    values = {
        "external_ref": _text(node, "ListingID"),
        "title": _text(node, "Title"),
        "portal_type": _text(details, "PropertyType"),
        "purpose": purpose,
        "description": _text(details, "Description"),
        "city": _text(location, ".//City"),
        "state": (state.get("abbreviation") if state is not None else "") or "",
        "neighborhood": _text(location, ".//Neighborhood"),
        "street": _text(location, ".//Address"),
        "number": _text(location, ".//StreetNumber"),
        "complement": _text(location, ".//Complement"),
        "zip_code": _text(location, ".//PostalCode"),
    }
    row.update({k: v for k, v in values.items() if v})
    numeric = {
        "sale_price": (details, "ListPrice"),
        "rent_price": (details, "RentalPrice"),
        "condo_fee": (details, "PropertyAdministrationFee"),
        "iptu": (details, "YearlyTax"),
        "area_total": (details, "LotArea"),
        "area_built": (details, "LivingArea"),
        "bedrooms": (details, "Bedrooms"),
        "bathrooms": (details, "Bathrooms"),
        "suites": (details, "Suites"),
        "parking_spaces": (details, "Garage"),
    }
    for field, (parent, tag) in numeric.items():
        text = _text(parent, tag)
        if not text:
            continue
        try:
            row[field] = parse_int(text) if field in INTEGERS else parse_decimal(text)
        except ValueError as exc:
            row["parse_errors"].append(f"{tag}: {exc}")
    row["features"] = [t.text.strip() for t in node.findall(".//Features/Feature") if t.text]
    row["photos"] = [
        i.text.strip()
        for i in node.findall(".//Media/Item")
        if i.text and (i.get("medium") or "image").lower() == "image"
    ]
    return row


def portal_types(rows):
    return sorted({r["portal_type"] for r in rows if r.get("portal_type")})


def apply_type_map(rows, type_map):
    """Preenche `category` das linhas do XML. Tipo ignorado ("") marca a linha como ignorada."""
    for row in rows:
        if "portal_type" not in row:
            continue
        target = type_map.get(row["portal_type"])
        row["skipped"] = target == ""
        if target:
            row["category"] = target
    return rows


# ---------- validação ----------


def check_rows(rows, existing_refs=()):
    """Valida as linhas sem gravar: devolve (novos, atualizados, ignorados); preenche `errors`."""
    seen, new, updated, skipped = set(), 0, 0, 0
    refs = set(existing_refs)
    categories = _category_names()
    for row in rows:
        row["errors"] = list(row.get("parse_errors", []))
        if row.get("skipped"):
            row["errors"] = []
            skipped += 1
            continue
        row["errors"] += _row_problems(row, seen, categories)
        ref = row.get("external_ref")
        if ref:
            seen.add(ref)
        if row["errors"]:
            continue
        if ref in refs:
            updated += 1
        else:
            new += 1
    return new, updated, skipped


def _category_names():
    """Nomes e prefixos de categoria em minúsculas (uma consulta para o arquivo todo)."""
    known = set()
    for name, prefix in Category.objects.values_list("name", "code_prefix"):
        known.add(name.lower())
        if prefix:
            known.add(prefix.lower())
    return known


def _row_problems(row, seen, categories):
    problems = []
    ref = row.get("external_ref")
    if not ref:
        problems.append("referência externa vazia")
    elif ref in seen:
        problems.append(f"referência repetida no arquivo: {ref}")
    if not row.get("title"):
        problems.append("título vazio")
    if not row.get("category"):
        problems.append("categoria vazia")
    elif str(row["category"]).strip().lower() not in categories:
        problems.append(f"categoria não encontrada: {row['category']}")
    if not (row.get("city") and row.get("neighborhood")):
        problems.append("cidade e bairro são obrigatórios")
    purpose = row.get("purpose") or "SALE"
    if purpose not in Property.Purpose.values:
        problems.append(f"finalidade inválida: {purpose}")
    if purpose in ("SALE", "BOTH") and not row.get("sale_price"):
        problems.append("preço de venda vazio")
    if purpose in ("RENT", "BOTH") and not row.get("rent_price"):
        problems.append("valor do aluguel vazio")
    if row.get("status") and row["status"] not in Property.Status.values:
        problems.append(f"situação inválida: {row['status']}")
    return problems


# ---------- ZIP de fotos ----------


class PhotoArchive:
    """Fotos de um ZIP lidas em memória (nada é extraído em disco)."""

    def __init__(self, fileobj):
        self.zip = zipfile.ZipFile(fileobj)
        infos = [i for i in self.zip.infolist() if not i.is_dir() and _is_image(i.filename)]
        if len(infos) > MAX_ZIP_MEMBERS or sum(i.file_size for i in infos) > MAX_ZIP_UNCOMPRESSED:
            raise ValidationError("ZIP grande demais.")
        self.by_name = {}
        self.by_folder = {}
        for info in infos:
            parts = info.filename.replace("\\", "/").split("/")
            self.by_name.setdefault(parts[-1].lower(), info)
            if len(parts) > 1:
                self.by_folder.setdefault(parts[-2].lower(), []).append(info)

    def photos_for(self, ref, names):
        """([(nome, bytes)], [nomes não encontrados]) dos arquivos citados ou da pasta `ref`."""
        if not names:
            infos = sorted(self.by_folder.get((ref or "").lower(), []), key=lambda i: i.filename)
            return [(i.filename, self.zip.read(i)) for i in infos if i.file_size], []
        found, missing = [], []
        for name in names:
            info = self.by_name.get(name.replace("\\", "/").split("/")[-1].lower())
            if info is None:
                missing.append(name)
            elif info.file_size:
                found.append((info.filename, self.zip.read(info)))
        return found, missing

    def close(self):
        self.zip.close()


def _is_image(name):
    return name.lower().endswith(IMAGE_EXTENSIONS) and not name.startswith("__MACOSX")


# ---------- modelo de planilha ----------


def sheet_template():
    """Bytes do XLSX modelo (cabeçalho + linha de exemplo)."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Imóveis"
    sheet.append([c[0] for c in COLUMNS])
    sheet.append([c[1] for c in COLUMNS])
    notes = book.create_sheet("Instruções")
    for line in (
        "ref_externa, titulo e categoria são obrigatórios; cidade e bairro também.",
        "finalidade: Venda, Aluguel ou Venda e aluguel. Informe o preço correspondente.",
        "situacao (opcional): Rascunho, Publicado, Reservado, Vendido, Alugado. Padrão: Rascunho.",
        "comodidades e fotos: separe por ponto e vírgula.",
        "fotos: nomes de arquivos do ZIP ou endereços http(s).",
        "Sem a coluna fotos, o ZIP pode ter uma pasta com o nome de cada ref_externa.",
        "Reenviar a mesma ref_externa atualiza o imóvel em vez de duplicar.",
    ):
        notes.append([line])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
